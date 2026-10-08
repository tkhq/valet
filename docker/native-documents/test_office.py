"""Exercise real save/reopen cycles, bounded preservation, and rendering."""
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen.canvas import Canvas

import office


class OfficeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def validate_unchanged(self, path):
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        self.assertTrue(office.inspect(path)['ok'])
        self.assertEqual(before, hashlib.sha256(path.read_bytes()).hexdigest())

    def docx(self):
        path = self.root / 'sample.docx'
        doc = Document()
        p = doc.add_paragraph()
        p.add_run('Original').bold = True
        p.add_run(' unchanged').italic = True
        doc.sections[0].header.paragraphs[0].text = 'Header retained'
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Table retained'
        doc.save(path)
        edited = Document(path)
        edited.paragraphs[0].runs[0].text = 'Edited'
        edited.save(path)
        actual = Document(path)
        self.assertEqual(actual.paragraphs[0].text, 'Edited unchanged')
        self.assertTrue(actual.paragraphs[0].runs[0].bold)
        self.assertTrue(actual.paragraphs[0].runs[1].italic)
        self.assertEqual(actual.sections[0].header.paragraphs[0].text, 'Header retained')
        self.assertEqual(actual.tables[0].cell(0, 0).text, 'Table retained')
        return path

    def xlsx(self):
        path = self.root / 'sample.xlsx'
        book = Workbook()
        sheet = book.active
        sheet['A1'] = 4
        sheet['B1'] = '=A1*2'
        sheet['B1'].font = Font(bold=True)
        sheet.merge_cells('A3:B3')
        sheet['A3'] = 'Merged retained'
        sheet.freeze_panes = 'A2'
        book.save(path)
        edited = load_workbook(path, data_only=False)
        edited.active['A1'] = 6
        edited.save(path)
        actual = load_workbook(path, data_only=False)
        self.assertEqual(actual.active['A1'].value, 6)
        self.assertEqual(actual.active['B1'].value, '=A1*2')
        self.assertTrue(actual.active['B1'].font.bold)
        self.assertIn('A3:B3', actual.active.merged_cells)
        self.assertEqual(actual.active.freeze_panes, 'A2')
        actual.close()
        return path

    def pptx(self):
        path = self.root / 'sample.pptx'
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        run = shape.text_frame.paragraphs[0].add_run()
        run.text, run.font.bold = 'Original', True
        slide.notes_slide.notes_text_frame.text = 'Notes retained'
        deck.save(path)
        edited = Presentation(path)
        edited.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].text = 'Edited'
        edited.save(path)
        actual = Presentation(path)
        self.assertEqual(actual.slides[0].shapes[0].text, 'Edited')
        self.assertTrue(actual.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.bold)
        self.assertEqual(actual.slides[0].shapes[0].left, Inches(1))
        self.assertEqual(actual.slides[0].notes_slide.notes_text_frame.text, 'Notes retained')
        return path

    def pdf(self):
        source, path = self.root / 'source.pdf', self.root / 'sample.pdf'
        canvas = Canvas(str(source))
        canvas.drawString(72, 720, 'Content retained')
        canvas.save()
        writer = PdfWriter(clone_from=str(source))
        writer.add_metadata({'/Title': 'Edited title'})
        writer.write(path)
        actual = PdfReader(path)
        self.assertEqual(actual.metadata.title, 'Edited title')
        self.assertIn('Content retained', actual.pages[0].extract_text())
        self.assertEqual(len(actual.pages), 1)
        return path

    def test_roundtrips(self):
        for create in (self.docx, self.xlsx, self.pptx, self.pdf):
            with self.subTest(format=create.__name__):
                self.validate_unchanged(create())

    def test_cli_invalid_file(self):
        path = self.root / 'broken.docx'
        path.write_bytes(b'not a zip')
        result = subprocess.run([sys.executable, str(Path(office.__file__)), 'validate', str(path)], capture_output=True, text=True)
        self.assertEqual(result.returncode, 1)
        self.assertFalse(json.loads(result.stdout)['ok'])

    @unittest.skipUnless(shutil.which('libreoffice') and shutil.which('pdftoppm'), 'Render tools require the sandbox image')
    def test_render(self):
        out = self.root / 'rendered'
        out.mkdir()
        for create in (self.docx, self.xlsx, self.pptx):
            path = create()
            # Separate directories avoid basename collisions between formats.
            target = out / path.suffix[1:]
            target.mkdir()
            profile = self.root / ('profile-' + path.suffix[1:])
            subprocess.run(['libreoffice', '-env:UserInstallation=' + profile.as_uri(), '--headless', '--convert-to', 'pdf', '--outdir', str(target), str(path)], check=True, capture_output=True, timeout=90)
            pdf = target / 'sample.pdf'
            self.assertGreater(len(PdfReader(pdf).pages), 0)
            subprocess.run(['pdftoppm', '-f', '1', '-singlefile', '-scale-to', '600', '-png', str(pdf), str(target / 'page')], check=True, capture_output=True, timeout=30)
            self.assertGreater((target / 'page.png').stat().st_size, 100)


if __name__ == '__main__':
    unittest.main()
