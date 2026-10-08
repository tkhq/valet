"""Read-only document inventory and structural checks. Never modifies input files."""
import argparse
import importlib.metadata
import json
from pathlib import Path
import shutil
import sys
import zipfile

from defusedxml.ElementTree import fromstring

PACKAGES = ('python-docx', 'openpyxl', 'python-pptx', 'pypdf', 'reportlab')
TOOLS = ('libreoffice', 'pdfinfo', 'pdftoppm', 'pdftotext')


def check():
    packages = {name: importlib.metadata.version(name) for name in PACKAGES}
    tools = {name: shutil.which(name) for name in TOOLS}
    return {'ok': all(tools.values()), 'packages': packages, 'tools': tools}


def inspect(path):
    path = Path(path)
    kind = path.suffix.lower()
    result = {'path': str(path), 'format': kind.lstrip('.'), 'bytes': path.stat().st_size}
    if kind in ('.docx', '.xlsx', '.pptx'):
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            # Bound decompression before reading untrusted packages.
            if len(entries) > 10000 or sum(e.file_size for e in entries) > 512 * 1024 * 1024:
                raise ValueError('Package exceeds inspection limits. Inspect it with a dedicated local tool.')
            names = archive.namelist()
            if len(set(names)) != len(names):
                raise ValueError('Duplicate package parts. Repair the source document before editing.')
            bad = archive.testzip()
            if bad:
                raise ValueError(f'Corrupt package part: {bad}. Recover an intact source document.')
            for name in names:
                if name.endswith(('.xml', '.rels')):
                    fromstring(archive.read(name))
            result['parts'] = names
            result['external_relationship_parts'] = [
                name for name in names if name.endswith('.rels') and any(
                    item.get('TargetMode') == 'External' for item in fromstring(archive.read(name))
                )
            ]
    if kind == '.docx':
        from docx import Document
        doc = Document(path)
        result.update(paragraphs=len(doc.paragraphs), tables=len(doc.tables), sections=len(doc.sections))
    elif kind == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False, keep_links=True)
        try:
            result['sheets'] = [{'name': s.title, 'rows': s.max_row, 'columns': s.max_column} for s in book]
        finally:
            book.close()
    elif kind == '.pptx':
        from pptx import Presentation
        deck = Presentation(path)
        result.update(slides=len(deck.slides), width=deck.slide_width, height=deck.slide_height)
    elif kind == '.pdf':
        from pypdf import PdfReader
        with path.open('rb') as stream:
            doc = PdfReader(stream, strict=True)
            if doc.is_encrypted:
                raise ValueError('Encrypted PDF. Supply an authorized decrypted copy before inspection.')
            result.update(pages=len(doc.pages), page_sizes=[
                [float(p.mediabox.width), float(p.mediabox.height)] for p in doc.pages
            ])
    else:
        raise ValueError('Unsupported format. Supply a DOCX, XLSX, PPTX, or PDF file.')
    result['ok'] = True
    result['scope'] = 'Readability and package structure only; verify content, formulas, and rendered layout separately.'
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest='command', required=True)
    commands.add_parser('check', help='Check installed libraries and render tools.')
    for command in ('inspect', 'validate'):
        sub = commands.add_parser(command, help='Read package structure without changing the file.')
        sub.add_argument('path')
    args = parser.parse_args()
    try:
        result = check() if args.command == 'check' else inspect(args.path)
        if not result['ok']:
            result['action'] = 'Use a sandbox image that includes the native document runtime.'
    except Exception as error:
        result = {'ok': False, 'error': str(error), 'action': 'Check the input file and runtime before retrying.'}
    print(json.dumps(result, indent=2))
    return 0 if result['ok'] else 1


if __name__ == '__main__':
    sys.exit(main())
